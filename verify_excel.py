import openpyxl

file_path = r"C:\Users\anand_fua08yg\.gemini\antigravity-ide\scratch\srcc_timetable_scraper\SRCC_Free_Classrooms_Timetable.xlsx"
wb = openpyxl.load_workbook(file_path, data_only=True)

print("Workbook Sheets and Row Counts:")
for name in wb.sheetnames:
    ws = wb[name]
    print(f" * Sheet '{name}': {ws.max_row} rows, {ws.max_column} cols")

ws_r = wb['Classrooms (R1-R37)']
print("\nSample R2 Data from Classrooms (R1-R37):")
# Find R2
for r in range(1, ws_r.max_row):
    v = ws_r.cell(row=r, column=1).value
    if v and 'ROOM NO: R2' in str(v):
        print(f"Found at row {r}:")
        for sub_r in range(r, r + 8):
            row_vals = [str(ws_r.cell(row=sub_r, column=c).value or '')[:40] for c in range(1, 6)]
            print(f"  Row {sub_r}: {row_vals}")
        break

ws_pb = wb['Practical Blocks (PB)']
print("\nSample PB3 Data from Practical Blocks (PB):")
# Find PB3
for r in range(1, ws_pb.max_row):
    v = ws_pb.cell(row=r, column=1).value
    if v and 'ROOM NO: PB3' in str(v):
        print(f"Found at row {r}:")
        for sub_r in range(r, r + 8):
            row_vals = [str(ws_pb.cell(row=sub_r, column=c).value or '')[:40] for c in range(1, 6)]
            print(f"  Row {sub_r}: {row_vals}")
        break
